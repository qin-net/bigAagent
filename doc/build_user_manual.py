"""Build doc/bigAagent-使用说明书.docx from live screenshots."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "doc" / "manual-screenshots"
OUT = ROOT / "doc" / "bigAagent-使用说明书.docx"
FONT = "宋体"
TITLE_FONT = "黑体"


def _east_asia(run, name: str) -> None:
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), name)


def run_text(paragraph, text, *, size=12, bold=False, font=FONT, color=None, center=False):
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    _east_asia(run, font)
    return run


def add_para(doc, text, *, size=12, bold=False, font=FONT, space_after=8, first_line=True, center=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.35
    if first_line and not center:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run_text(p, text, size=size, bold=bold, font=font, center=center)
    return p


def add_heading_cn(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.first_line_indent = Cm(0)
    size = 16 if level == 1 else 14 if level == 2 else 12
    run_text(p, text, size=size, bold=True, font=TITLE_FONT)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.first_line_indent = Cm(0)
    run_text(p, text, size=10.5, center=True, color=RGBColor(0x44, 0x44, 0x44))
    return p


def add_picture(doc, name, caption, width=15.2):
    path = SHOTS / name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width))
    add_caption(doc, caption)


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (left, right) in enumerate(rows):
        cell0, cell1 = table.rows[i].cells
        cell0.text = ""
        cell1.text = ""
        p0 = cell0.paragraphs[0]
        p1 = cell1.paragraphs[0]
        run_text(p0, left, size=10.5, bold=True, font=TITLE_FONT)
        run_text(p1, right, size=10.5)
        p0.paragraph_format.first_line_indent = Cm(0)
        p1.paragraph_format.first_line_indent = Cm(0)
    doc.add_paragraph()
    return table


def set_doc_defaults(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE


def build() -> None:
    doc = Document()
    set_doc_defaults(doc)

    for _ in range(4):
        doc.add_paragraph()
    add_para(doc, "bigAagent", size=22, bold=True, font=TITLE_FONT, first_line=False, center=True, space_after=6)
    add_para(doc, "私人投研模拟器  使用说明书", size=18, bold=True, font=TITLE_FONT, first_line=False, center=True, space_after=18)
    add_para(doc, "A 股研究辅助工作台  V0.1.0", size=14, font=TITLE_FONT, first_line=False, center=True, space_after=8)
    add_para(doc, "延迟行情 · 研究辅助 · 非投资建议", size=12, first_line=False, center=True, space_after=24)
    add_para(doc, "所有界面截图均来自本软件实际运行状态。示例标的为五粮液（000858），行情采集时间为 2026年9月4日 14:47。", size=10.5, first_line=False, center=True)

    doc.add_page_break()
    add_heading_cn(doc, "目  录", 1)
    for line in [
        "第一部分  软件简介 ................................................ 3",
        "第二部分  安装说明 ................................................ 5",
        "第三部分  使用手册 ................................................ 7",
    ]:
        add_para(doc, line, first_line=False, space_after=4)

    doc.add_page_break()
    add_heading_cn(doc, "第一部分  软件简介", 1)
    add_para(
        doc,
        "bigAagent V0.1.0 面向 A 股研究辅助需求，提供延迟行情浏览、四维首次分析、模拟账户记账、手动追踪复核、用户口径沉淀与专家私有记忆展示。用户在本机启动 Web 服务后，用浏览器打开工作台，按“分析—投入—追踪—沉淀”完成一次完整研究链路。学费停在模拟盘，不连接实盘交易通道。",
    )
    add_para(
        doc,
        "系统里只有五位业务 Agent：基本面、技术面、情绪、宏观与追踪。首次研究由应用层并行叫醒四位分析专家，再由系统做综合决策；追踪由追踪 Agent 对照原判断决定维持、复核或证伪。综合决策、编排、校验与落盘是 Python 系统环节，不包装成第六个 Agent。对外口径始终是研究辅助，不是投资建议。",
    )

    add_heading_cn(doc, "1.1  软件名称", 2)
    add_para(doc, "中文：bigAagent 私人投研模拟器 / A 股研究工作台", first_line=False)
    add_para(doc, "英文：bigAagent", first_line=False)
    add_para(doc, "简称：bA / bigAagent  V0.1.0", first_line=False)

    add_heading_cn(doc, "1.2  软件适用行业与用途", 2)
    add_para(doc, "适用领域：A 股研究辅助、投研流程演练、多 Agent 协作实验、模拟组合记账及相关个人学习。", first_line=False)
    add_para(doc, "用途：在延迟行情上并行完成基本面、技术面、情绪与宏观分析，由用户决定是否用虚拟资金投入，并在需要时手动触发追踪；记住的口径须过闸门，专家只继承自己的私有记忆。", first_line=False)

    add_heading_cn(doc, "1.3  软件开发平台", 2)
    add_para(doc, "软件运行平台：macOS、Windows、Linux（本机浏览器访问本地服务）。", first_line=False)
    add_para(doc, "编程语言：Python 3.9 及以上；前端为静态 HTML / JavaScript / CSS。", first_line=False)
    add_para(doc, "版本号：V0.1.0", first_line=False)
    add_para(doc, "程序量：约 18953 行源代码（src 下 Python、JavaScript、CSS、HTML）。", first_line=False)

    add_heading_cn(doc, "1.4  创作目的", 2)
    add_para(doc, "（1）把研究、决策、追踪和纪律留在同一条链路上，避免“看完研报另开交易软件、下次分析忘掉上次主张”。")
    add_para(doc, "（2）用四位分权专家交卷，主张必须对上当轮快照；追踪对照原判断，而不是每天重写一份荐股。")
    add_para(doc, "（3）用户口径不落原话、不当行情证据；专家记忆按角色隔离；知识库短卡片须人批准后才可被检索。")
    add_para(doc, "（4）虚拟资金按延迟价成交，把试错成本停在模拟盘。")

    add_heading_cn(doc, "1.5  系统总流程图", 2)
    add_para(doc, "一次典型路径为：启动本地服务并打开封面 → 进入工作台并采集延迟行情 → 搜索标的并打开详情 → 四专家并行完成首次分析 → 用户决定是否模拟投入 → 需要时手动追踪 → 在用户画像与专家工作台查看沉淀。", first_line=False)
    add_picture(doc, "fig-flow.png", "图1-1  bigAagent 一次研究主路径")

    add_heading_cn(doc, "1.6  主要功能", 2)
    add_para(doc, "（1）封面与工作台导航模块")
    add_para(doc, "启动后首先显示项目封面，说明五位专家职责与“非投资建议”口径。点击“进入工作台”后，可通过封面、市场、自选、模拟账户、用户画像、专家工作台六个入口切换。")
    add_para(doc, "（2）延迟行情采集与市场浏览模块")
    add_para(doc, "点击“采集行情”将全市场延迟报价写入本地库。市场页可按代码或名称查询、分页浏览，并打开单标的详情。页眉会提示数据可能过期及最近一次成功采集时间。")
    add_para(doc, "（3）四维首次分析模块")
    add_para(doc, "在标的详情中可补充本轮口径（可空），点击“开始首次分析”。基本面、技术面、情绪、宏观四位专家并行交卷，系统给出综合决策。完成后显示四维结论、重要留意、已记住口径以及各专家是否沿用上次记忆。")
    add_para(doc, "（4）反馈与按维重跑模块")
    add_para(doc, "对结论不满意时，可在反馈框写要求并“提交并按需重跑”。带“记住”且过闸门的短规则会留下；点名某一维可只重跑该维。用户原话不存盘，不作为行情证据。")
    add_para(doc, "（5）模拟投资与选股记忆模块")
    add_para(doc, "首次分析完成后解锁投入。填写数量与买入理由，按当前延迟价用虚拟资金成交。持仓随行情重估。买入理由进入选股记忆；也可只加入自选、不买。")
    add_para(doc, "（6）手动追踪模块")
    add_para(doc, "有研究基线后，由用户选择时间触发追踪。追踪 Agent 对照原判断输出维持、需要复核或原判断失效，并可指出继续盯紧的条件。没有按小时自动巡检。")
    add_para(doc, "（7）用户画像与专家工作台模块")
    add_para(doc, "用户画像由模型根据已记住口径、选股记忆与模拟行为归纳，不存原话，可停用口径并重新生成画像。专家工作台按角色展示跨标的经验；“分次分析记录”默认收起。追踪不能改写四位分析专家的 State。")

    add_heading_cn(doc, "第二部分  安装说明", 1)
    add_heading_cn(doc, "2.1  软件运行环境", 2)
    add_para(doc, "本软件为本地 Web 应用：后端 FastAPI 服务与研究队列在本机运行，前端由浏览器打开。分析、追踪与画像生成需要模型密钥；仅浏览已入库行情不需要密钥。运行环境见表2-1。", first_line=False)
    add_para(doc, "表2-1  运行环境", first_line=False, center=True, space_after=6)
    add_table(
        doc,
        [
            ("最低硬件配置", "64 位处理器，8GB 及以上内存，2GB 以上可用磁盘；可用集成显卡。"),
            ("推荐硬件配置", "Apple Silicon 或同级 x86 四核以上，16GB 内存，固态硬盘。"),
            ("客户端操作系统", "macOS、Windows 10/11 或常见 Linux 发行版。"),
            ("Python 环境", "Python 3.9 及以上，建议使用仓库内虚拟环境 .venv。"),
            ("支撑软件", "现代浏览器（Chrome / Safari / Edge 等）；看板 extra 含 FastAPI、Uvicorn、AkShare。"),
            ("模型服务", "分析 / 追踪 / 画像需配置 DEEPSEEK_API_KEY；默认模型为 deepseek-v4-flash。"),
        ],
    )

    add_heading_cn(doc, "2.2  安装过程说明", 2)
    add_para(doc, "（1）取得源码并检查目录")
    add_para(doc, "将仓库完整放到本地目录，保持 pyproject.toml、src 下运行时与工作台代码、.env.example 等结构不变。关键文件如图2-1所示。")
    add_picture(doc, "fig-install-dir.png", "图2-1  仓库关键文件与启动入口")

    add_para(doc, "（2）创建虚拟环境并安装依赖", first_line=False)
    add_para(doc, "在仓库根目录执行：", first_line=False)
    add_para(doc, "python3 -m venv .venv", first_line=False, font="Menlo")
    add_para(doc, ".venv/bin/python -m pip install --upgrade pip", first_line=False, font="Menlo")
    add_para(doc, '.venv/bin/python -m pip install -e ".[dev,board,akshare]"', first_line=False, font="Menlo")

    add_para(doc, "（3）配置环境变量", first_line=False)
    add_para(doc, "复制 .env.example 为 .env，填写 DEEPSEEK_API_KEY。已经在 shell 中 export 过的变量不会被覆盖。改完后必须重启 serve。示例文件如图2-2所示（图中密钥栏为空，请勿把真实密钥写入说明书或截图）。")
    add_picture(doc, "fig-env.png", "图2-2  环境变量示例（.env.example）")

    add_para(doc, "（4）启动方式", first_line=False)
    add_para(doc, "在仓库根目录执行 .venv/bin/python -m insightboard serve（当前代码入口）。一条命令同时拉起接口、静态页面和研究后台。浏览器打开 http://127.0.0.1:8000 ，应看到封面，如图2-3所示。改完 Python 路由后需重启服务；仅改页面脚本刷新浏览器即可。")
    add_picture(doc, "fig-cover.png", "图2-3  启动成功后的项目封面")

    add_heading_cn(doc, "2.3  备注", 2)
    add_para(doc, "（1）行情为延迟数据，页眉会提示最近成功采集时间；过期数据仍可用于复盘界面操作，但不代表盘中实时价格。")
    add_para(doc, "（2）未配置模型密钥时，仍可采集行情、浏览市场和模拟账户；开始分析、追踪或生成画像会提示无法调用模型。")
    add_para(doc, "（3）研究任务由后台串行执行。同一时刻若已有任务进行中，新请求会提示已有任务并继续轮询，而不是把错误显示成无意义对象。")
    add_para(doc, "（4）本地数据默认写在 data/ 下（行情库、研究日志库、方法论库等），分析日志不进入 git。")
    add_para(doc, "（5）本软件不提供实盘下单。虚拟盈亏不能当作投资建议或收益承诺。")

    add_heading_cn(doc, "第三部分  使用手册", 1)
    add_para(
        doc,
        "本部分按照“封面—市场行情—打开标的—首次分析—模拟投入—手动追踪—学习沉淀—K 线与公告—自选—模拟账户—用户画像—专家工作台”的顺序，完整展示一次以五粮液（000858）为例的研究任务。所有界面截图均来自本软件 V0.1.0 实际运行状态。",
        first_line=False,
    )

    add_heading_cn(doc, "3.1  打开封面并进入工作台", 2)
    add_para(doc, "浏览器访问 http://127.0.0.1:8000 后进入封面。确认标题为“私人投研模拟器”，并看到分析、投入、追踪、沉淀四步以及五位专家职责。点击“进入工作台”。页脚口径为“延迟行情 · 研究辅助 · 非投资建议”，如图3-1所示。")
    add_picture(doc, "fig-cover.png", "图3-1  项目封面")

    add_heading_cn(doc, "3.2  采集并浏览市场行情", 2)
    add_para(doc, "进入工作台后默认停在“市场”。页眉显示数据状态；本说明书截取时提示“数据可能过期，最近成功：2026/9/4 14:47:31”。需要更新时点击“采集行情”，等待全市场延迟报价入库。表格给出代码、名称、市场、最新价、涨跌幅和成交额，如图3-2所示。")
    add_picture(doc, "fig-market.png", "图3-2  市场行情页")

    add_heading_cn(doc, "3.3  按代码查询标的", 2)
    add_para(doc, "在搜索框输入股票代码或名称，点击“查询”。以 000858 为例，结果仅保留五粮液一行：深市、最新价 71.89、涨跌幅 1.71%、成交额 31.02 亿，如图3-3所示。点击该行进入详情。")
    add_picture(doc, "fig-market-search.png", "图3-3  查询 000858 五粮液")

    add_heading_cn(doc, "3.4  打开标的详情与决策流程", 2)
    add_para(doc, "详情页上方显示当前研究标的名称、代码、延迟价和涨跌幅，并提供加入自选、用户画像、模拟账户入口。中部“从研究到持续决策”给出四步轨道：首次分析、投资决策、持续追踪、沉淀迭代。首次分析完成前，投入与追踪步骤保持锁定。顶部上下文如图3-4所示。")
    add_picture(doc, "fig-detail-hero.png", "图3-4  标的详情与决策轨道")

    add_heading_cn(doc, "3.5  执行首次分析并阅读四维结论", 2)
    add_para(doc, "在“给本轮分析补充口径”中可留空，也可写一句可执行规则（例如“记住：估值须核对经营现金流”）。点击“开始首次分析”，状态从排队/研究中变为研究完成。本例在 9/4 14:49 完成：综合决策为持有，信心 56%；基本面持有、技术面卖出、情绪持有、宏观暂不判断。各位专家卡片会标明是否沿用上次记忆。可在底部反馈框点“提交并按需重跑”，如图3-5所示。")
    add_picture(doc, "fig-analyze.png", "图3-5  AI 首次分析（四维结论）")

    add_heading_cn(doc, "3.6  选择是否模拟投入", 2)
    add_para(doc, "分析完成后，“选择是否投资”解锁。用户先阅读结论，再决定数量。本例五粮液已投入 1100 股，平均成本 71.46 元，市值约 7.91 万元。填写数量（默认 100 股）和买入理由后点“确认投入”；理由进入选股记忆。也可“减少持仓”或保持观察不买，如图3-6所示。")
    add_picture(doc, "fig-invest.png", "图3-6  模拟投资决策")

    add_heading_cn(doc, "3.7  手动触发追踪", 2)
    add_para(doc, "有基线之后，在“本轮追踪重点”中可写关注点（可空），点击“开始本轮追踪”。本例 9/4 14:52 的追踪策略为“维持原判断”，摘要为无实质变化、维持持有，并列出相对基线的变化与继续盯紧项（经营现金流、关键价位、量能）。追踪后同样可以反馈并按需重跑，如图3-7所示。")
    add_picture(doc, "fig-track.png", "图3-7  手动追踪结果")

    add_heading_cn(doc, "3.8  查看本标的学习沉淀", 2)
    add_para(doc, "“Agent 迭代与用户画像”汇总本标的记忆沉淀，并可跳转完整用户画像。分析、反馈和追踪过程中会积累专家私有记忆与用户偏好，不把用户原话当作研究证据，如图3-8所示。")
    add_picture(doc, "fig-learn.png", "图3-8  本标的学习沉淀")

    add_heading_cn(doc, "3.9  查看行情走势", 2)
    add_para(doc, "详情下方“行情走势”提供收盘折线与日 K，并可“拉取日线”。走势仅作研究辅助，不构成买卖时点建议，如图3-9所示。")
    add_picture(doc, "fig-chart.png", "图3-9  行情走势")

    add_heading_cn(doc, "3.10  查看最新公告", 2)
    add_para(doc, "“最新公告”列出已采集的信息披露标题与时间。拉取日线时会尝试一并采集。公告用于情绪维对照事实，不把标题当结论，如图3-10所示。")
    add_picture(doc, "fig-notices.png", "图3-10  最新公告")

    add_heading_cn(doc, "3.11  管理自选股", 2)
    add_para(doc, "导航点击“自选”，或在详情中“加入自选”。自选页列出关注标的与最新延迟价格，点击代码可回到详情，如图3-11所示。")
    add_picture(doc, "fig-watchlist.png", "图3-11  自选股")

    add_heading_cn(doc, "3.12  查看模拟账户", 2)
    add_para(doc, "“模拟账户”显示总资产、现金、持仓市值与相对本金。本例总资产 994,295.00 元，相对本金 −0.57%。持仓表按延迟价重估浮动盈亏；下方为最近成交与选股记忆。页内声明虚拟资金模拟、非投资建议，如图3-12所示。")
    add_picture(doc, "fig-paper.png", "图3-12  模拟账户")

    add_heading_cn(doc, "3.13  查看并生成用户画像", 2)
    add_para(doc, "“用户画像”由模型基于选股策略与行为证据生成，不存原话。本例标题为“现金流校验的价值型投资者”，风险倾向显示为杠铃型，并列出偏好策略与决策方式。页面下部可查看、停用已记住的研究口径，并可“重新生成画像”，如图3-13所示。")
    add_picture(doc, "fig-profile.png", "图3-13  我的投资者画像")

    add_heading_cn(doc, "3.14  查看专家工作台", 2)
    add_para(doc, "“专家工作台”按专家归类决策偏好，并展示跨任务经验总结。每位专家只继承自己的私有记忆；用户口径会约束下一次怎么看。数字“覆盖 N 只标的”是最新记忆覆盖的股票数，不是追踪把四专家全员重跑的次数。“分次分析记录”默认收起。追踪不能改写分析专家 State，如图3-14所示。")
    add_picture(doc, "fig-experts.png", "图3-14  专家工作台")

    add_heading_cn(doc, "3.15  常见问题及处理方法", 2)
    add_para(doc, "浏览器打不开页面", first_line=False, bold=True, font=TITLE_FONT)
    add_para(doc, "确认已在仓库根目录执行 python -m insightboard serve，并用 http://127.0.0.1:8000 访问；改端口时加上 --port。")
    add_para(doc, "提示未配置模型密钥", first_line=False, bold=True, font=TITLE_FONT)
    add_para(doc, "检查 .env 中的 DEEPSEEK_API_KEY，保存后重启 serve。仅采集行情不需要密钥。")
    add_para(doc, "采集行情失败", first_line=False, bold=True, font=TITLE_FONT)
    add_para(doc, "系统会保留上一批成功数据。可稍后重试，或运行 python -m insightboard status 查看最近采集。")
    add_para(doc, "开始分析提示已有任务进行中", first_line=False, bold=True, font=TITLE_FONT)
    add_para(doc, "研究队列串行执行。页面会继续轮询已有任务；不要同时对多只股票连点分析。长时间无结果可重启 serve，后台会回收超时任务。")
    add_para(doc, "追踪按钮不可用", first_line=False, bold=True, font=TITLE_FONT)
    add_para(doc, "须先有成功的首次分析基线。分析失败或仍在进行中时，追踪保持锁定。")
    add_para(doc, "记住口径没有留下", first_line=False, bold=True, font=TITLE_FONT)
    add_para(doc, "记住须是可执行的短规则（例如带“对照”“核对”），不能是情绪发泄或事后涨跌总结。原话不会存盘。")
    add_para(doc, "把结论当成买卖指令", first_line=False, bold=True, font=TITLE_FONT)
    add_para(doc, "四维结论与模拟盈亏都是研究辅助。延迟行情、虚拟资金，不构成投资建议，也不等于实盘可成交价格。")

    add_para(
        doc,
        "软件流程小结：启动服务并打开封面，采集延迟行情后选股；完成首次分析并阅读四维结论，再决定是否模拟投入；需要时手动追踪；在画像与专家工作台管理口径与记忆。全程保持“研究辅助、非投资建议”。",
        space_after=12,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build()
